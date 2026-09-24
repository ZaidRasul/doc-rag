import os
import tempfile
from pathlib import Path
from typing import List
from langchain_core.documents import Document
from langchain_community.document_loaders import PyMuPDFLoader,TextLoader,UnstructuredMarkdownLoader, Docx2txtLoader
from docx import Document as DocxDocument
supported_extensions = [".pdf", ".txt", ".md", ".docx"]

def load_files(uploaded_files) -> List[Document]:
    all_docs = []
    # Load multiple files if needed
    for uploaded_file in uploaded_files:
        file_name = uploaded_file.name
        file_bytes = uploaded_file.getvalue()
        extension = Path(file_name).suffix.lower()
        # check if the file extension is supported
        if extension not in supported_extensions:
            raise ValueError(f"Unsupported file type {extension}")
        # check extensions and call the appropriate loader function
        if extension == ".pdf":
            documents = load_pdf(file_name, file_bytes)
        if extension == ".docx":
            documents = load_docx(file_name, file_bytes)
        if extension == ".txt":
            documents = load_txt(file_name, file_bytes)
        if extension == ".md":
            documents = load_md(file_name, file_bytes)

        all_docs.extend(documents)

    return all_docs

def load_pdf(file_name: str, file_bytes: bytes) -> List[Document]:
    # Using PyMuPDFLoader to load PDF files
    # But it needs a file path, so we use tempfile
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=".pdf"
            ) as temporary_file:
                temporary_file.write(file_bytes)
                temp_path = temporary_file.name
        loader = PyMuPDFLoader(temp_path)
        documents = loader.load()

        for document in documents:
            page_number = document.metadata.get("page", 0)
            document.metadata = {
                "source": file_name,
                "file_type": "pdf",
                "page": page_number,
            }
        return documents
    finally:
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)


def load_txt(file_name: str, file_bytes: bytes) -> List[Document]:
    # we could directly decode the bytes to string so we don't need to use TextLoader
    # as text files dont need parsing, we can just read the content as is
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("utf-8", errors="replace")
    document = Document(
        page_content=text,
        metadata={
            "source": file_name,
            "file_type": "txt",
        },
    )
    return [document]



def load_md(file_name: str, file_bytes: bytes) -> List[Document]:
    # we could directly decode the bytes to string so we don't need to use UnstructuredMarkdownLoader
    # as md files dont need parsing, we can just read the content as is
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("utf-8", errors="replace")

    return [
        Document(
            page_content=text,
            metadata={
                "source": file_name,
                "file_type": "md",
            },
        )
    ]



def load_docx(file_name: str, file_bytes: bytes) -> List[Document]:
    #2 ways to do it here, either use python-docx or langchain's docx2txt.
    # Using python-docx to load DOCX files
    temp_path = None
    try: 
        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=".docx"
        ) as temporary_file:
            temporary_file.write(file_bytes)
            temp_path = temporary_file.name

        docx_doc = DocxDocument(temp_path)
        text = "\n".join([paragraph.text for paragraph in docx_doc.paragraphs])

        return [
            Document(
                page_content=text,
                metadata={
                    "source": file_name,
                    "file_type": "docx",
                },
            )
        ]
    finally:
            if temp_path and os.path.exists(temp_path):
                os.remove(temp_path)
    